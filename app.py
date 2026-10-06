# -*- coding: utf-8 -*-
"""
Ranking de ventas de feria — Kbas Office
Sube el Excel con las líneas de venta (referencia + cantidad, y precio si hay)
y un zip con las fotos (nombre = referencia). Devuelve el ranking de lo más y
menos vendido, con foto, en PDF y/o Excel, separado por marca (Kbas / Volum).
"""

import io, os, re, zipfile, tempfile
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image as PILImage

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph,
                                Image as RLImage, PageBreak)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# ---------- Paleta ----------
# Paleta oficial Kbas
INK="1A1A1A"; BURD="522003"; SAND="E1C9AF"; MUTE="7A6A5A"  # ink · chocolate(acento) · sand · texto suave
LOGO_B64 = "iVBORw0KGgoAAAANSUhEUgAAAKEAAADICAYAAACAqoEkAAAj5UlEQVR4nO19eZgkVZXv73cis6q6m266qumuiqikCrFEaEAEQRzFQQQF/HyCoIL78kR0RnH9HHX0ofPGbRhFZhy3UWdkxIVBHfA5KouCKIIKyCCLAi3VXRWR1QW9SHdTS8Y974+82WblUpkZkVmZUZW/76uvOyPucuLec88999xzz6XrugcBeD/JEwH0qqoDQACApAOAABxVJQCHpNjfDgAp+i32PQE4AMTmkXxRdJBH4d+GQBLGmOuDIHg+AG0kr+u6HyB5sqoeSXKtpUlVNbRJlGSoqiSpAArPYdMU12cKJBWS2P+zKE2hXQpwbFmFtrGfRLH5QgA7SN6tqt/0ff9bjXxff3//gX19fb8jmWkkXxEU9ruKvteQNPhzOxmSoU1nip4ZAKGqGgCh/W3sb2M/tJDvMVX9cX9//6X33nvvXKHyFID3ici7LAHI81A5Sp/X+l3tWVSoKgAcNzIyMrR169agwbxnOo5zkjFmwfOl/oYaZR1I8gnGmIcANMSEPT09x5IcjkMa7OAoEhZ/fllHu9TbliRP2blz5/0Ari48FwCHAfs7udOxdm5urr/RTCTnk/B9lsZco/kcxxliM0dLi1DoA1U9tPi5WFGZFCganIpXCBLVJiTni39LtYRddNEqWL17P6RIOe+ii7ZAkDBR3kU5jDEN65FtxkJJ2C4quli5UNUFgyZpTKgiYmon6yJJSNp03F0dL0MIyaTpE12UwO5SJBZJm467qIwuEy4h1HGcrklpmSFRTEhSSw2dXSQPpSpgopiwi8oQkUTr9V0mXB5I9OyQNCY0c3NzXZ1wmSFpTNjFMkTXgWEZIIF7xwsgSLiNqYvkI2nTsenuHS8/JI0Jyxwiu0g+EseEXSwLdP0Ju2gvujsmXXQcukzYRduRNKfWLpYhuk6tXbQd3em4i7YjaUxIG1SoiwTDGLNgqzhxTGiMSRrNXZSjzE64ErbBuk4aHYwV4UXT3errbHSntmUAEUm1m4ZGUBoJrmsn7DCsBKldFpWrayfsOCx79agU3em4iyWHMWZhkMxSm00XXSw1JGFxTKSnpydK9P/uQOsglC1MVohOuKyVfXudR5KQXGO1qq6EHZOGB4y9QyYxKNu2K42a2eFY9kwYpT+SJglLVUBBsvQlMcYkqsGXCIlpE3th0wKeSyXpHhOSkkqllrUkjAKSyd4xSdh0LGEYphvNtNz3x0Wk4TZpJ8qm44StjmmMidLgy3p1bIzpaTcNjSAMw4Wn7ZLGhCQT1eBLhL52E9AIRGThdGyM+VO7iIkCEeltNw2tRBQdnWSSmDAkubf4gZB8WJNwBSbyV5Wq6up209FiRLHbrmk6FS2Cqu6em5vbXvxMUqnUvQAeaRNNUbDcmbBhqGoimNDed/zQ9u3bp4ufy9atWwNV/VUCrssFAKjq2nbT0IFITJuo6rUokfYCACQvbwtFESAi69pNQ6eB5IHtpqEeGGNmjDHfLH0uAJBOp69R1TuSIA2NMevbTUOnQVU7fmBa3vr21NTUPaXvBADGx8dnALw3CYZrkkPtpqHDQAAHtZuIxWAXlFlVvbjS+/1bYL7v3wDg7xIgDZ/cbgI6CSMjI+sBjHaygUNVc6r6lmw2O17p/YJ9WN/3P2qM+ZdOZURVhao+ZWRkxG03LZ2CMAyfSnKw3XQsAgPg7UEQ/Fe1BKXOACYIgotU9ROdyogi0p/L5V7cbjo6BcaYV7EDO8uStBfAG33f/9xiaSt5pBjf998fhuFrAGzvtO+z0867RkZG+ttNS7vhed5TSZ7XaVOx1QHvMca8wPf9f6uVvqpbVDab/Q+SJxljriwU3Ckg+cT5+flPt5uOdqK/v/9AAF8g2TGGassj+4wxnxaRv8xmsz+rJ9+ivnmTk5MPBEFwnjHmTFX9aVFFbYWqQkRe53neP42NjS3rveRKGBwc3LRq1apvkTyxE6Sg5YlZY8y3AJwUBMG7JyYmdtSbvy4H0Ww2+yPf909V1eer6lUAHrNbMNGobgKsh+7b9u7d+9+ZTObpi6VNktOnqi42qOh53oscx7mR5BntZMCi/p8C8HlVfWYQBC/3ff/OhsuKQsDw8PCTVPVsAOcAOK7gXtWORiEJY8wsgGtIfl1VbwmCoHgvPOW67u9E5MmdIDUWg9Wlvuv7/rnFzz3POxjAc1X1tSRPsenaQh8AqOpjAG5R1f/M5XI/mJ6ezsYqNyZdMjw8fLQx5gySZyLPkGstoTGLbgxFDZQFcDeAO1XVB/AMkucvKTExoKqzJD9PcquqHgrgWFU9UkTW2/dLRkvxTKeq06p6q4j8QESu37Zt20NNq6dZBQHAyMjIoWEYnqSqpwP4CwBPKGKOZla1KErVhE6XgKVoJ/1F/ZUD8HtVvRnAj3t6em4dHx+PJfGq1tmKQgFgw4YNa3t6eo4h+VwApwI4tl1SsovFUcR40wB+TfI6Vb1x9erV9z344IOzLa+/1RUUMDg4+IRUKvVsVT0LwOkk13SZsb2wuuU0gGtU9fthGN66ffv2qSWnY6krBICDDz74iWEYfgTAK9pFQxeYJfmPqnqp7/uPtpOQJT/Dm8lkhnO53EtJnoAuA7YTKVU9DcCLrOG7bVgKJuDg4OAhjuOcqKpnkDyT5KbuVNx+FOmCDwH4fySvS6VSt7dqAVKVjlYUau1ax6rqX5J8FoAju4uSzkYRQz6iqneRvJnkzWEY3p3NZqdrZI9XdzMKGRwc3ETyOBF5DoCTkGe69YX3XcZLFkrMalMA7iJ5kzHmZzMzM3fv3Llzd1Pri5pxaGjoCJKnkTwdwAkkNxXetXs7qRhJHACVtkM7oU0tDeMAbiH5Q2PMTUEQbI1dfiOJBwcHD0mlUmer6jmqeryIrCoirm0oaqS9qvoHAL8nmQVwAIDnAxhpI3kNQ1XvAvALexD+YACHA3giyXS72xpY0N67VPXnInKlMeaHJdul9ZdXT6KhoaHjSf41ybNI9lsCotTXdFhb1+2q+u+O4/xoYmLijygKdzc4OHiIiPyc5HAbyawL9ltuDsPwzKmpqf1RCsbGxnr37dt3BICzALyG5KGd1P4AYIyZIHlFKpX60tatW7c0VMZiLzOZzJiqflBVzyfZ2ykfDiw4PPPBNWvWfH0xy77neVeRPLeT6K8E64zxviAIPlktTSaTGQjD8CKSf4MOi0Fj+2QHgC+IyKfqdeeqaiccHh6+0BhzC4DXAuhEBvwfVT0lCIKv1NpaSsIpwgJqxaKZmJjYEQTBh+3O03Qn+HcWYHlkgOQHjDG/GBoaOqOefGVM6Lruas/zvgzgCwA2dhLzFaCq42EYnh0Ewf31pE9SJK96aQ2C4FoALzfGPN5ikhqG9fU8XESucV33fbXSL2DCjRs3HgDg2yT/dycyH5C/korke6ampv5Ybx5VTZJTa920+r5/A8nPdJI0LMDyT1pEPu667iWLpS1mwlQqlfpXEXlhBzMgjDF3TU5O/leD+RITYrjRIOgkP2eM2dUicmLDHsV4j+u6H6iWZn/neJ73DhE5v1MZsAg3AGhUx0sMEzaKycnJCQB3dqI0LMBOzx9xXff5ld4LAHiedxiADyWAAQGgLj1wJYHkH9pNQy2QTJG81Kp8C1CQEO8h2fFBdQBAVSfaTUMHYlu7CagFKw03p9PpV5W+E8/zDrY7IO2grSGoqiG5vXbKsnzLOno/AL/dBDSACwAsCH4vJF8gIhvaRFCjmEmn07si5Ov8EWbBaIHsW+rl0ixYQfcUz/OOLH4uxphT20NS41DVx2dmZvbWTrkQETu2LahlrK4Eu4fbCnKaDnsG/BnFz4TkUUn5AACP53K5ho2zSdoxiXL/NMk9SNZFmUcV/xZV3VQtcSfBmiAeX7t27VyEvIlhQpLztVOVYZ+qRsnXFpBcENpPSCYplsvsIYccEuUWzMR0UBRajTGzEZm3XVhwA0OijLgk52666aYo01ViOkhVG5b06XR6PkkqB0q8txLFhADmEe1S6oY7tl0QkYYHzNzcXA7JujK47Ob3xCCqvS9JTKiqDUc8EBHDkpszOxml/ShJMuRGMV/YfIlhwiiqg+M4BglaHZdCkCBDLiLSKiItj6fSLESUhIoEMWGptUKQIF1iJUzHIpIYWmOg7NLtxIygqIgiXdoBzV+RkZiVfFSU6YRIkCSMgUQwIQAYYxqWhLt371ZN0LYXSnguUQuTqEiSnTCKicbqhIlhwjKdMElbWjGQGJVjJahHZUy4EiShMSYxUiIMw2XPhKW7O4LGz2skDiQTw4RI0LQaAytyOk5Kx2rCBkwklEnClTAdr5CBlhiULhRXxHSM5EjCFYEySdiVEh0FXamr4y4TdhYaltpr166VhEWZKJeEnXx6v4vlh5VqJ1z235gwJHfvmGRDwYKK0F2YdBDKHBiiOoq2CVH1hqToGytlsCx05Sp9sBwhIkmJT0hVbXiBkcvlBEDUWWLJUSr4VoRO2GjMvzZCojBhApHog06RmMmGnkgEHMdJDK1RUUknTJKdMJJup6rp2qk6A8aYhplQVYnk6L1lSNpBp0gguayZ0BiTtGm8zJUrMQuTqLsCSZKEIpIYWqOi1FNIVDUxTIiIOmxSJCHJSCv5BE7HyV2YREWSJGEUWvv6+gQJYsJKntVJ0gkjNXZSJCGQLFqbhe7quMMQ5eIfY0yivGhKkSjCre7TMJIkXYwxDdOaQJ1wARLFhMhvTTVMc5Ik4UpYHZciaUwYCUmShFEGTBiGkQZnpyBRhJOUk08+ueFpJ0mSMEkDpllIFBMC4J49e5b16jjKgEm8ThhV2W8Tlv3qOMqAsVt2SerHBRAAifHaUFVn9+7dDUvv5S4Jk46kTceRkKSOjTJg0um0kyQHhlLXOolxbqMdkPn5+Sg6YWKkfZKkdgwsvEIiQV7HAMAwDKOsjhPDhBEXJonSCUv7I1E6IaI3dJKky7JfHZfOvpIU8W8vbZYokjBJ07Gq9kTIkxgGBMqlfdI8cqM2dsMd2y6QjMKETsIcGBb0YwrArSRPBfZfipxP1YbQIIvF/hYRhGG4KwiCKFcs7ExCqBOSUNWGL9Du6enZEYbhDMkDFmvDdvdpoX5V/W1xmtTMzMwlfX19vQBOI7kO1sewQuiMVlzYUroootUXFPnRIqpKkkZVfQAfQYRQdmEYfghATlUzWDgKS7+pEKRy0WOwNlJ+aAl2sLiEpl38laYp/faQ5L0kP79Y3ZWwbds23/O896rqhQD6bZst+A5VNSUR/peiP4vbJ1TVR1T1W9ls9roFaYoLcF231xhDEVHHcRYMKcdxNJ1ON9UBdn5+fsFqNwxDGmP2/7aHuiEiOj09vacJVZYyQ7Wo983snGqLhtLnzbgaLLVhw4ZVQL7NbFR/AEA6nV5Q9lL0p30mhfonJiYavjC9iy666KKLLrrooosuuuiiiy666KKLLrrooosuuuhiZaDzd/VXAIaHhzMAxlT1UFU9BEAGwCaSa60DKEnOAXhEVW8HcAeA+4Ig2IaljyXkIL+92LR66XneC0k+xd4JvH9/0caoKTgSaHEkJRv42qjqvKreks1m7y0u1HXd1QAuILlJVcOCE6Mto3B1lrG/DcnQhpAt/H/GGHPD1NTUwxs3bjwglUpdRvIwVd1Fcpeq7iE5o6rzfyaJvdZPrTCwlGTKPit82wE2HW2eEMAeVd1O8gGSvzXG/E8QBPua1cDV4HneU1X1bADPI7kZwHrrRQMA21V1N4BVIpKplN++/72q3gTgmiAIfokmXgcyMDCwbvXq1YeHYXgMySMBHKKqQwBWA8iRfAzAjKrOkfRV9R4A9zuOc8/ExMRkI3XR87zHSB4QhVDbaFdOTk6eV/zcdd2TROTmKGUWyjXGfNP3/VcMDQ2d7DjOjVHKUdW9qvpKABMikjPGbALwOREZq5JeVfUPJL8bhuGXpqamHo76DdUwNDR0hoi8HcBzC76DRc4tvjHm7DAM75uenn48k8n0qOozjDGXkyxjxiLXKAVwC4BP+r7//ai0DQwMrOvp6XmeiLwYwEkARljk/1UYJNbBGEWDZj+MMTssLdc4jvPTiYmJB2vVm0Le1Sgq3dXOb6yKU6Yt9wAAsJKr4fy2ge4KguDq4uee5/03gIuqlEkReTKA9zuO80bXdf8+CIJ/RhOmHs/zDgPwcQDnFHdmMb3GmDuy2eyvC8+s18lPPc+7jOQlpTQX/SbJZ6nqNZ7nfcb3/fcCWHCd62IYGhra6DjOBar6epJjxWUXGM4O6CtIXgPAD8NwSETOB/AqFJ3aJDlA8oUAXmiM2eN53m9IXp3L5a6empr6Y6X6U4jfwGX5RSSMy4SF2y7jXPZT5cbMVYvlKaJ7o4hc5nne8ar65jhTtOd5LwPwz1Y9qeq8W82r2hhzj+Msfh6tUCbJd7iuuyoIgregjr71PO91AC4GcEhxOUU0QVUfJvnyycnJW0uy/9B13WtJfhlAXyktyKs/zwHwHMdxLvY87z9JXjI5OflAcSGC1pw9bqayvLrJHsFr6k1opcCrSX5tbGysN0plruu+HcA3AGyqY2AOoPJise72VFWIyIWu67510YoGBtZ5nvc1kv9G8pBqtKnqjKq+pgIDAgCCILgCwEer9VHRoFtP8gJVvcXzvNcXp+n4MCAtOK7ZEDNZRnzJ3r17/2+jFbmu+0aSl6L+25ZCVGA4x3HWNlKvpfnvMplMRd13dHR0fW9v73dIvqaGZIaq/iQIgkX1+3Q6/VljTM3FiK3nIJJf9TzvrwrPk3A4pqku6KrasESznfouz/NOrTfP8PDwM0hehsbMYDOVHhpj1jdQBgCA5HpjzPsqvErNz89/RUROq0dlUtW7aqUZHx/fRfK2emcsy/j/ODw8fAzQIiY0xjQtBLGqRjnYVBVRj3+SdFT1Q6ijzUZHR/uMMZeRXN1IHVr9ireGyrFlAcB5mzZtOrT4ued5byJ5Tr06O8n1dVbZkFlGRFap6iuAfIO2IgJDM3RCBQDHcWbjLnIWFBpxerfS8Fme5z2lVtrZ2dmXiMjTG6W72iKMZN0r3ZJ8BziO85LCb8/zNqjq+xuhi+Szn/a0p9U8m66qOxulT1WfDUA6eToudEjH3LNijd9n1kjmiMibI5Y/W+m5qu6KUp7FGUX/f1E143eVekHyqCAI3lQrrTVe1w1b9tjIyMhgJzNhqxB3IXbSYi9d132Sqh4fRXoXnzQshohEUkksDYcNDAyss79rDaCKZZD8pOu6Z9VIF0Va94dhOLISmTAunjQ6OtpX7aWqHi8ikcw5qK77RdZHSPb39PT05//LI6KUoaprSF7hed75i6SJwkspVW2KJOxoE08FRO5QK1kGZmZm1i2SLFJHA4uGAInTxj0A1oyOjvaq6voY+vUaAJcPDw+/o9JLEWnIjFSAMaYndkixFtjx9hfdonJjgWQvyaq7LiIyWCVfzT9VfbRS3pht7DiOk3rsscd6SEaV0AWkAVzqed7XhoaGNpa821PPNxb/WeK2pRgzkA4r3AjlOI5jTLz1xCLminbDOI5TlTbr6aPWtPS4/dsHYK74myzT7SGZBfCIMWaS5DdbQK+S1L6+vrSqpuLuPlkd8TUicoLneW/zff8GAJiZmflyb2/vHICT7PZkH/PeUyk7VYulJaeqsyKyT1V/uX79+jubIcXKuC2ifrAAVfZ9Y6PSoGkEqrqnr6+v6kowl8t9OJVKfYPkXsdxHiP5uDFmZnBwMHf77beXSveluNJNc7lcDvkwgE1ZA1hGPEJVf+C67sVBEPzDjh07/gTgs/avGCz6A0p8EYMgiB8gs4rEaoae2HGS0EqRiS1btlRlQhsz547S5xMTE5HrVdXHY0gwIyJGRMIwDJvWpla37BWRT3ied4Ix5kPZbPa+SklRQ7VqhsQq+7Aot5cvFZowzf8KS2y7XEwHrSe7McYJw3AGVbYF48BKxXNF5FbP8y7zPO/gRstoBrOUdWrhfraYOx2tiuAUiyhVvaZZhJSAg4ODq0muTqfT63K53HoRWaeqGwF8IEZbMpVKORMTE497nreLpNfMHShgfz+vI3mRqr7U87yP+r7/RdSpbrRkOg7DsLeW/1sdGLfld4RUtU6nv12zZs3P4pa1YcOGtb29vUcBOFZVNwN4AslB5F251oZhuJrkKs3HGYw7mAvhhBX5Nt0cl/5F6gEAl+RnPc87B8AHfd//Za18zejgMmVFRBazo9WEtb7/IE4ZrYCq/v2DDz5YcWutFjZv3tyzc+fOUwC8jOQpyDPeguiplZityfvmd5BseNckQj0g+VxV/annef+eTqc/PD4+nq2Wvhk6YSVGflKM8gDgUt/3fxu1jGbDSsHLs9nsdxvNOzY21ut53ut27dr1S5I/EpE3AHgC8GeHz8V8+pqAQvRZiMgNLaxnYaV24ULywvn5+Z+4rnt4tbSxmbDUP2/z5s09qvq/Gi3HMl9ojLnE9/2/jUtXs2AZ8IbZ2dm3oUF90nXdk/bu3fsT6718HFCXZAutjXGfqu6I4p1SAp2fnzcAYIy5TfMHuWIW2UDl1pxD8rsHHXSQWylNM+xG+1duAwMD63bu3PkZETm6nhFXNB2FxpjrAZwRBMF70Vr7WV09ULSL8R+rVq0619rB6obnee8mea2IPLMeSWfruhPA8QCOyeVyR4dhuDmdTh+pqp+LyjiqqiISAoA9J/PVSAXFQIER0+n05ZlMpmylH0sntIUfPTw8/GpjzGEAXioiT16swQuNqao5Vb1fVX8kIlf5vn9bHFqqVVf8o7+//0CST140w59Pwt2GvFrw7UbrdF33UyTf2eg0S/LuSmrI0NDQlar6VxWy1ANTYEJbx5eNMX9NsmFTShxo/uzLacaYywBciKJZpRkLkxEAl4t1TVys0e3e6G0kr1fVm9asWXNPVEW/FiwdT3Jd95XWKXREVV9JcqwKjXsAPGaM+bWqfiqbzUZaBbuu+xEReWdE3aslDsbz8/P7mdD3/Uc9z7uY5FeXSj/cT0heaF0wNDR0XzabvbTwvCnmj3qnXgBXzMzMfKjRqS0GDgJwvtWr+gHsMMY8QHLBwknzUQS+p6r3kLwPQOi67nFhGE5u3759qt7KXNc9i+TfRu3cauaoah7X9YCkKd0C9X3/cs/zziJ5VjsYUUQ+5rrur4Mg+DnQJCYsmmIXrRzARb29ved4nvcbzYevuLmnp+ee8fHxplvy7bR6fxAECxZJhbAiIvKGAr3WherVJSYTTaVSj3qedyfJ7+Ryue9NTU1tr1bf6Ojo+rm5uU+RlDiG5SrP43CKEZHSHZ6Q5FuNMU8TkcxSMyLJPgCXjY6OPmt8fHymGavjWWPMNap6rarWvDmJZIbk2SJyKcnb5ufn7/A879NDQ0MnxKWlAmZQ0rHT09N7RORLWqXlix4T+eOJzwPwBcdx7vA87z2ospibnZ19vYg8sRUdWqzTNQq7MCnbZpycnJwA8FZVXfLjE3ZaPm52dvYlQMzVsWW464MgOMv3/dPDMDxeVb9RixGLFHbHLt/fKSK/8DzvqsXsSY1Cq9x+SXIf6tz/LaJ1mOQlruv+C0rabXR0tE9E3tgEklshCffbCUsRBMHVqvrBpTTZFENELgDgNMNYvV+/2759+xbf91+rqjfW+2FFnZwmeS7JG4eHh0+JS1eNOiN9t9Vn3ux53ruLn8/Ozp4I4Ii4UpBtuAA9CIKPq+q7scReS7bfT8hkMpsFMT1CKnRoLqpzphXTg8aYr2cymeE4dNWoJ7JHuWW0iwcHB48sPBORF7Bd4qQGSNIYs+ig833/06r6WgB7l/IzRGRVGIanSjXdqAGU5TfGjEcuLC9tPGPMG+KRBaD8LrtCHbEuqSa5xnGc9xd+AnhW1LJaDc17QdTs4yAIrlDVczQfq3EpSCvgmc2Yjst2N0RkVxMU9GfHLQAtOoRlv+3FmzZtOnTDhg0HIB9ltRVVFRDnO8zc3FxdO1BBEFyrqi9Q1YeWghFtmx3ejNVx2Qeq6iziT/ORAncuFUiudhznRalU6iAAG5pUbMX+MMYsma4YBMHtYRieZoz5NvKRc1taH0m3GZKw7GB2GIY5xPc+7vgz0SRPFhEPzTP6N/2bSYaVTDSLYWpq6uEgCM5X1TNU9dZWMqKqrm2GJCw7eS8iuWbZn1TVidEILZ0jVfVQkk1bQLXqcFdUBEFwXTqdPkVV3wugJdurJHuaMfLKJKFtzE44N1wx3l8zYFfyB6rqgU0stqLupvmb4yNBVZ1aq+PFMD4+PuP7/iUAXgrgkRZIxWhh0haUUGE6FhEThmHcUR0C+1eykVBtMBQ9j3vwP82IEbOqlFetzZrS85lM5mhVfZmqunafuhdYcFNDgQ4lOaf5QEzjxpgf+L7/fdd1zwBwOcnNTVyIhbFjVldamDQDzZiaLG2VmFCbYJoC8nuw2zV/zmbJDc11giR1YGBgnTHmGpKHlB4rqJBh/7/271wApwZBcPvg4OApjuN8DMDr0Ry9fSa2sRqtc0CNbcGvZD5qMvap6kMA9rayEifeqTGS1N7eXg95d7ayYwWL/dlIGke5rjsAAFNTU9t933+jqp4DwI87Pavqnzp2BarNCQPSMr3UNv6E67pbAEw0Q1eqtm0XUyWR/OzKVYguuXpFZEEksiAIrjbGnGaMuTPmtz/ajNVxWRRPa9dqeoybCGitBZm87fbbb59X1V81ozxVrRaVKw6oqnQcZw4RZz1VDSvF37ERF56vqj+Owoh2qv99M+yENUPJRkFB12T+mrGoZbRsU94uIgrHUptyID6O8+oikHQ6nZ6bm5tW1YaiqRZhH8mKwQiCIHikr6/vPFX9TURGvK5VkrAZwXcKTBgnQGRLFgvWYfZXvu//CgDm5uauV9U/NkE/qrbSjlVwGIbO1NTUNMkHG6XRSqs/rVmzpuplQlu2bNltjPlQo4s9VZ0Ske+3ShLGVpDiMF8z6VgEn4AdKI8++uhjKI9G1TBY5WaBmAZ7IN8OCuB7EfPvvvfeexcNWayqd2gDx1PtQP76tm3bfInb2ZWY0HGcFOIv3zsuKhew/xzylUEQLJiC5+fnv2SMuSsms1QMFxxHJSlGOp3+ijEmiEBjzd0Sx3HCRvR4VZ0G8E9AnlHiOhq0RCdEZ+y4LIAdvb8Lw/AilNBnjw1cGEPvgqpWC5/SlLYYHx/PknxHhC3VTbWukbDbl/11lpczxrw5CIKtQHPCgJQRF8dptAgtMx9F8Se0DHgfgHOrncCbnJy8zRjzBpL7IirpFadjEYljKQiLz6j4vn+lqr4LeUN7zcxWAo/5vn9ijXTPrGeRauv8SHFIlWYsTCo1XGxdLM5+aR1l1z1ICrsGqvojAKf7vv+HxdJns9mrjDEvjrJQqbSad133IGPMy6NO86paFt44CILLjDHnqerWesq1g+MfRkdH11dJklLV19UoA8hf0v1/fN//6ILMiG/PKwvGbY2j8TRpW64tK2oxFaexerYEC3Wq6sOq+qkgCL6IOu8QDoLg2oMPPvikMAwvBvBa1nFns31/hOu6H0b+lGCG+YP6J4rI+hgxwM3s7GwlG993MpnMrcaYv7E0rqtGo/V2/4v5+fkfDg8Pf4DkrfYu5vTw8PARqvo+kidWym8HMIwx15O82Pf9W0rTpFR1p4g0fHdaUQVlW1Yk9yB/1iSSvmiV/0kg7yArES6esrQdZi/xe7SwAAvDcK2qXiAiVTlbVXdpPi7Mlap6VRAEjzRa/7Zt23wAFw4NDf2riLxJVc9k/rjrYjQfSPJi22k7VPV+AF8Ow3Cc5FtI1h1b0Hr5AMBsGIYVzSsTExOTAC7yPO+zxpgLSL684JpWylC2vGeo6g3GmD94njetqgcaYw4Tkf2DrPj7jDGPquoNAL4SBMF1qCYUhoaGNtvzsqkGtocoIo6IzM/Nzd1SQUfi4ODgySLSsK+d5DnucWPMjUEQPOK67moAf0PyaAB7VTVnJUsaeb2mR/MRFHLIT7H7RYaqzpOc0nzUeldVe0iuBbAJgFFVQ/JxVd0tIr6qPqCqvxORe+253KYhk8kMqOoxAJ5qjBlDPjrEamug3gtgGsA2VX3IcZyHenp6tm7ZsmV3If/GjRuHUqnUu0gepaoDyN9PAuSlc45k2qpGYgf/GpI0xnwxCIKP1UOj67oHAThdRM5S1acjf8x1v7pVidGKofloYltU9ZcArnMc52eW0RdFR54QayFY8v+aQb07GAWdudgtTQBwbGzMmZ+fX6Wq3Lp1a6TQcplMZsAYcziAYwEcA+AwVR0G0G/3oQvHOB4h+QDJXwD4SS6Xu3tqaqohh47/D/ciq5pZILvjAAAAAElFTkSuQmCC"

# ---------- Detección flexible de columnas ----------
REF_CANDS   = ["REFERENCIA","REF.","REF","ARTICULO","ARTÍCULO","CODIGO","CÓDIGO","COD"]
QTY_CANDS   = ["CANTIDAD VENDIDA","CANTIDAD","CANT.","CANT","UDS","UNIDADES","QTY","VENDIDO"]
PRICE_CANDS = ["PRECIO","PVP","P.V.P","IMPORTE UNIT"]
DTO_CANDS   = ["DTO","DTO.","DESCUENTO","% DTO","DESC"]

def _norm(s): return re.sub(r"\s+", " ", str(s)).strip().upper()

def _find_col(cols, cands):
    m = {_norm(c): c for c in cols}
    for cand in cands:
        if _norm(cand) in m:
            return m[_norm(cand)]
    return None

def limpiar_ref(v):
    # quita TODO espacio/invisible (en cualquier posición) y pasa a mayúsculas
    return re.sub(r"\s+", "", str(v)).upper()

def leer_excel(file_like):
    """Busca la pestaña y columnas de referencia + cantidad (prefiere 'LINEAS')."""
    xls = pd.ExcelFile(file_like)
    orden = sorted(xls.sheet_names, key=lambda n: 0 if "LINEA" in _norm(n) else 1)
    for name in orden:
        df = xls.parse(name)
        rc = _find_col(df.columns, REF_CANDS)
        qc = _find_col(df.columns, QTY_CANDS)
        if rc and qc:
            return df, rc, qc, _find_col(df.columns, PRICE_CANDS), _find_col(df.columns, DTO_CANDS), name
    raise ValueError("No encuentro columnas de Referencia y Cantidad en ninguna pestaña del Excel.")

def construir_ranking(df, rc, qc, pc, dc, ordenar_por="unidades"):
    d = df.copy()
    d = d[d[rc].notna()]                        # fuera filas sin referencia (incluye la de TOTAL)
    d["_REF"] = d[rc].map(limpiar_ref)
    d = d[d["_REF"] != ""]
    d["_Q"] = pd.to_numeric(d[qc], errors="coerce").fillna(0)
    if pc:
        precio = pd.to_numeric(d[pc], errors="coerce").fillna(0)
        dto = pd.to_numeric(d[dc], errors="coerce").fillna(0) if dc else 0
        d["_IMP"] = precio * (1 - dto/100) * d["_Q"]
    else:
        d["_IMP"] = 0.0
    d["_MARCA"] = d["_REF"].str[:1].map(lambda x: "Volum" if x == "V" else "Kbas")
    out = {}
    for marca, sub in d.groupby("_MARCA"):
        g = (sub.groupby("_REF").agg(uds=("_Q","sum"), euros=("_IMP","sum"))
                 .reset_index().rename(columns={"_REF": "ref"}))
        col = "euros" if ordenar_por == "euros" else "uds"
        g = g.sort_values(col, ascending=False)
        out[marca] = [{"ref": r.ref, "uds": int(r.uds), "euros": round(float(r.euros), 2)} for r in g.itertuples()]
    return out

def _es_imagen(nombre):
    return nombre.lower().endswith((".jpg", ".jpeg", ".png"))

def recoger_fotos(archivos, tmpdir):
    """Acepta una lista de archivos subidos: imágenes sueltas y/o zips (mezcla).
    Devuelve {REF: ruta_foto}. Ignora ocultos del Mac."""
    m = {}
    for f in archivos:
        nombre = f.name
        datos = f.getvalue()
        if nombre.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(datos)) as z:
                    for n in z.namelist():
                        base = os.path.basename(n)
                        if not base or base.startswith("._") or "__MACOSX" in n or not _es_imagen(base):
                            continue
                        dst = os.path.join(tmpdir, base)
                        with z.open(n) as src, open(dst, "wb") as out:
                            out.write(src.read())
                        m[os.path.splitext(base)[0].upper()] = dst
            except zipfile.BadZipFile:
                continue
        elif _es_imagen(nombre):
            base = os.path.basename(nombre)
            if base.startswith("._"):
                continue
            dst = os.path.join(tmpdir, base)
            with open(dst, "wb") as out:
                out.write(datos)
            m[os.path.splitext(base)[0].upper()] = dst
    return m

def hacer_thumb(src, tmpdir, ref):
    try:
        im = PILImage.open(src).convert("RGB")
        im.thumbnail((130, 180))
        dst = os.path.join(tmpdir, f"thumb_{ref}.jpg")
        im.save(dst, quality=85)
        return dst
    except Exception:
        return None

def fmt_eur(v): return f"{v:,.2f} €".replace(",", "§").replace(".", ",").replace("§", ".")
def fmt_num(v): return f"{int(v):,}".replace(",", ".")

# ---------- Generar EXCEL ----------
def generar_excel(ranking, fotos, tmpdir, con_euros):
    wb = Workbook(); wb.remove(wb.active)
    thin = Side(style="thin", color="D9D2C7")
    for marca in ["Kbas", "Volum"]:
        if marca not in ranking: continue
        ws = wb.create_sheet(marca)
        ws.merge_cells("A1:D1")
        ws["A1"] = f"Ranking de ventas · {marca}"
        ws["A1"].font = Font(name="Arial", bold=True, size=13, color=INK)
        ws.row_dimensions[1].height = 24
        heads = ["Foto", "Referencia", "Cantidad vendida", "€ vendidos"]
        for c, h in enumerate(heads, 1):
            cell = ws.cell(row=2, column=c, value=h)
            cell.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=BURD)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions["A"].width = 11
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 14
        ws.freeze_panes = "A3"
        r = 3
        for it in ranking[marca]:
            ws.row_dimensions[r].height = 56
            src = fotos.get(it["ref"])
            if src:
                th = hacer_thumb(src, tmpdir, it["ref"])
                if th:
                    img = XLImage(th); img.width = 52; img.height = 72
                    ws.add_image(img, f"A{r}")
            ws.cell(row=r, column=2, value=it["ref"]).font = Font(name="Arial", size=10)
            cc = ws.cell(row=r, column=3, value=it["uds"]); cc.font = Font(name="Arial", size=10); cc.alignment = Alignment(horizontal="center", vertical="center")
            cd = ws.cell(row=r, column=4, value=it["euros"] if con_euros else None)
            cd.font = Font(name="Arial", size=10); cd.number_format = '#,##0.00 €'; cd.alignment = Alignment(horizontal="right", vertical="center")
            for col in range(1, 5): ws.cell(row=r, column=col).border = Border(bottom=thin)
            r += 1
        ws.cell(row=r, column=2, value="TOTAL").font = Font(name="Arial", bold=True, size=10, color=INK)
        tc = ws.cell(row=r, column=3, value=f"=SUM(C3:C{r-1})"); tc.font = Font(name="Arial", bold=True, size=10); tc.alignment = Alignment(horizontal="center")
        if con_euros:
            td = ws.cell(row=r, column=4, value=f"=SUM(D3:D{r-1})"); td.font = Font(name="Arial", bold=True, size=10); td.number_format = '#,##0.00 €'; td.alignment = Alignment(horizontal="right")
        for col in range(1, 5): ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=SAND)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0); return buf

# ---------- Generar PDF ----------
def generar_pdf(ranking, fotos, tmpdir, con_euros, titulo):
    styles = getSampleStyleSheet()
    h_st   = ParagraphStyle("h", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=16, textColor=colors.HexColor("#"+INK), spaceAfter=2)
    sub_st = ParagraphStyle("s", parent=styles["Normal"], fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#"+MUTE), spaceAfter=8)
    cell_st= ParagraphStyle("c", parent=styles["Normal"], fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#"+INK))
    refb_st= ParagraphStyle("r", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor("#"+INK))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=14*mm, bottomMargin=14*mm,
                            leftMargin=16*mm, rightMargin=16*mm, title=titulo)
    story = []
    marcas = [m for m in ["Kbas", "Volum"] if m in ranking]
    for mi, marca in enumerate(marcas):
        items = ranking[marca]
        tot_u = sum(x["uds"] for x in items); tot_e = sum(x["euros"] for x in items)
        story.append(Paragraph(titulo, h_st))
        resumen = f"{marca} · {len(items)} referencias · {fmt_num(tot_u)} uds"
        if con_euros: resumen += f" · {fmt_eur(tot_e)}"
        story.append(Paragraph(resumen, sub_st))
        data = [[Paragraph("<b>Foto</b>", cell_st), Paragraph("<b>Referencia</b>", cell_st),
                 Paragraph("<b>Cantidad</b>", cell_st), Paragraph("<b>€ vendidos</b>", cell_st)]]
        for it in items:
            src = fotos.get(it["ref"])
            th = hacer_thumb(src, tmpdir, it["ref"]) if src else None
            if th:
                im = RLImage(th); im.drawHeight = 14*mm; im.drawWidth = 14*mm*0.806
            else:
                im = Paragraph("<font color='#B0A99E'>— sin foto —</font>", cell_st)
            euro = fmt_eur(it["euros"]) if con_euros else "—"
            data.append([im, Paragraph(it["ref"], refb_st), Paragraph(fmt_num(it["uds"]), cell_st), Paragraph(euro, cell_st)])
        data.append([Paragraph("", cell_st), Paragraph("<b>TOTAL</b>", refb_st),
                     Paragraph(f"<b>{fmt_num(tot_u)}</b>", cell_st),
                     Paragraph(f"<b>{fmt_eur(tot_e) if con_euros else '—'}</b>", cell_st)])
        t = Table(data, colWidths=[20*mm, 60*mm, 35*mm, 43*mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#"+BURD)),
            ("TEXTCOLOR",  (0,0), (-1,0), colors.white),
            ("ALIGN", (2,0), (3,-1), "RIGHT"), ("ALIGN", (0,0), (0,-1), "CENTER"),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0,1), (-1,-2), [colors.white, colors.HexColor("#F6F2EC")]),
            ("BACKGROUND", (0,-1), (-1,-1), colors.HexColor("#"+SAND)),
            ("LINEBELOW", (0,0), (-1,0), 0.5, colors.HexColor("#"+BURD)),
            ("TOPPADDING", (0,1), (-1,-1), 3), ("BOTTOMPADDING", (0,1), (-1,-1), 3),
            ("LEFTPADDING", (0,0), (-1,-1), 5),
        ]))
        story.append(t)
        if mi < len(marcas) - 1: story.append(PageBreak())
    doc.build(story); buf.seek(0); return buf

# ==================== INTERFAZ ====================
def check_password():
    def entered():
        if st.session_state.get("password") == st.secrets.get("dashboard_password"):
            st.session_state["ok"] = True; del st.session_state["password"]
        else:
            st.session_state["ok"] = False
    if st.session_state.get("ok"): return True
    st.title("Ranking de ventas")
    st.text_input("Contraseña", type="password", on_change=entered, key="password")
    if st.session_state.get("ok") is False: st.error("Contraseña incorrecta.")
    return False

st.set_page_config(page_title="Ranking de ventas · Kbas Office", page_icon="🏆")
if not ("dashboard_password" in st.secrets and not check_password()):

    st.markdown("""<style>
      .stApp { background:#FDF9F3; }
      div.stButton > button { background:#522003; color:#FDF9F3; border:0; border-radius:8px; font-weight:600; }
      div.stButton > button:hover { background:#6B2B06; color:#ffffff; }
      div.stDownloadButton > button { background:#A88EA1; color:#1A1A1A; border:0; border-radius:8px; font-weight:600; }
      div.stDownloadButton > button:hover { background:#916F88; color:#ffffff; }
    </style>""", unsafe_allow_html=True)
    st.markdown(
        '<div style="text-align:center;padding:6px 0 2px">'
        '<img src="data:image/png;base64,' + LOGO_B64 + '" style="height:64px">'
        '<div style="margin-top:6px"><span style="font-family:monospace;font-size:12px;letter-spacing:.14em;'
        'background:#1A1A1A;color:#E1C9AF;padding:3px 10px;border-radius:6px;">OFFICE</span></div>'
        '<h1 style="color:#522003;margin:12px 0 2px;font-size:30px;font-weight:600;">Ranking de ventas</h1></div>',
        unsafe_allow_html=True)
    st.caption("Sube el Excel con las líneas (referencia + cantidad, y precio si lo hay) y las fotos "
               "(sueltas o en ZIP). Devuelve el ranking de lo más y menos vendido, con foto, separado por marca.")

    c1, c2 = st.columns(2)
    with c1:
        excel_file = st.file_uploader("1 · Excel con las líneas de venta", type=["xlsx", "xlsm"])
    with c2:
        fotos_files = st.file_uploader("2 · Fotos (sueltas o en ZIP · nombre = referencia)",
                                       type=["zip", "jpg", "jpeg", "png"], accept_multiple_files=True)

    o1, o2 = st.columns(2)
    with o1:
        ordenar = st.radio("Ordenar por", ["Unidades", "Euros"], horizontal=True)
    with o2:
        formatos = st.multiselect("Formato de salida", ["PDF", "Excel"], default=["PDF", "Excel"])
    titulo = st.text_input("Título del documento", value="Ranking de ventas")

    if excel_file and fotos_files and formatos:
        if st.button("Generar ranking", type="primary"):
            try:
                df, rc, qc, pc, dc, hoja = leer_excel(excel_file)
            except Exception as e:
                st.error(str(e)); st.stop()
            ranking = construir_ranking(df, rc, qc, pc, dc, "euros" if ordenar == "Euros" else "unidades")
            con_euros = pc is not None

            with tempfile.TemporaryDirectory() as tmp:
                fotos = recoger_fotos(fotos_files, tmp)
                # resumen de cobertura
                tot_ref = sum(len(v) for v in ranking.values())
                con_foto = sum(1 for v in ranking.values() for it in v if it["ref"] in fotos)
                st.success(f"Procesado: {tot_ref} referencias · {con_foto} con foto "
                           f"({round(100*con_foto/tot_ref) if tot_ref else 0}%). "
                           f"Columnas usadas: Ref='{rc}', Cantidad='{qc}'"
                           + (f", Precio='{pc}'" if pc else " (sin precio → sin euros)")
                           + (f", Dto='{dc}'" if dc else "") + f". Hoja: '{hoja}'.")
                for marca, items in ranking.items():
                    st.write(f"**{marca}**: {len(items)} referencias, {sum(i['uds'] for i in items)} uds")

                if "PDF" in formatos:
                    pdf = generar_pdf(ranking, fotos, tmp, con_euros, titulo)
                    st.download_button("⬇ Descargar PDF", pdf, file_name="ranking.pdf", mime="application/pdf")
                if "Excel" in formatos:
                    xls = generar_excel(ranking, fotos, tmp, con_euros)
                    st.download_button("⬇ Descargar Excel", xls, file_name="ranking.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.info("Sube el Excel y las fotos (sueltas o en ZIP), elige el formato y pulsa Generar.")
